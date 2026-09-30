"""Bouwt een PowerPoint uit de metingen van meet.py, de GIF's van gif.py en de grafieken uit grafieken.py.

    python bouw.py <deckmap met project/> <werkmap> <decksleutel> <uitvoer.pptx>

Tekst, vormen, lijnen en tabellen worden echte PowerPoint-objecten op de gemeten plaats. Canvasgrafieken
worden een echte grafiek met een wipe-animatie (als grafieken.REGISTER ze kent) of een GIF die vanzelf speelt.
Sprekersnotities komen uit <aside>; elke dia krijgt een fade-overgang zoals in het webdeck.
"""
import json
import math
import re
import sys
from pathlib import Path
from urllib.parse import unquote, urlparse

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

import grafieken

DECK, WERK, SLEUTEL, UIT = Path(sys.argv[1]), Path(sys.argv[2]), sys.argv[3], Path(sys.argv[4])
EMU = 6350
FONTS = {'DM Serif Display', 'DM Sans', 'JetBrains Mono', 'Segoe UI', 'Georgia', 'Arial'}


def e(v):
    return Emu(int(round(v * EMU)))


def rgb(k):
    return RGBColor.from_string(k['hex'].upper())


def alfa(kleur_el, a):
    """Zet doorzichtigheid op een a:srgbClr."""
    if a is not None and a < 0.999:
        el = kleur_el.makeelement(qn('a:alpha'), {'val': str(int(a * 100000))})
        kleur_el.append(el)


def vul(vorm, k, op=1.0):
    if not k:
        vorm.fill.background()
        return
    vorm.fill.solid()
    vorm.fill.fore_color.rgb = rgb(k)
    alfa(vorm.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr')), k['a'] * op)


def lijn(vorm, k=None, breedte=0):
    if not k or breedte <= 0:
        vorm.line.fill.background()
        return
    vorm.line.color.rgb = rgb(k)
    vorm.line.width = Pt(breedte / 2)


# ------------------------------------------------------------------ tekst
def paragrafen(runs):
    """Splitst runs op harde regeleinden en vouwt witruimte samen zoals CSS dat doet."""
    par, huidig = [], []
    for r in runs:
        stukken = r['t'].split('\n')
        for i, s in enumerate(stukken):
            if i:
                par.append(huidig)
                huidig = []
            if s:
                huidig.append(dict(r, t=re.sub(r'\s+', ' ', s)))
    par.append(huidig)
    for p in par:
        if p:
            p[0]['t'] = p[0]['t'].lstrip()
            p[-1]['t'] = p[-1]['t'].rstrip()
    return [[r for r in p if r['t']] for p in par if any(r['t'].strip() for r in p)]


def lettertype(r):
    f = r['f'] if r['f'] in FONTS else ('DM Sans' if 'sans' in r['f'].lower() else r['f'])
    if f == 'DM Sans' and r['b'] == 500:
        return 'DM Sans Medium', False
    return f, r['b'] >= 600


UITLIJNING = {'left': PP_ALIGN.LEFT, 'start': PP_ALIGN.LEFT, 'right': PP_ALIGN.RIGHT, 'end': PP_ALIGN.RIGHT,
              'center': PP_ALIGN.CENTER, 'justify': PP_ALIGN.JUSTIFY}


def vul_tekst(tf, runs, al, lh, eerste_grootte=None):
    pars = paragrafen(runs)
    for i, p in enumerate(pars):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = UITLIJNING.get(al, PP_ALIGN.LEFT)
        grootte = max((r['s'] for r in p), default=eerste_grootte or 24)
        if lh:
            para.line_spacing = Pt(lh / 2)
        for r in p:
            run = para.add_run()
            run.text = r['t'].upper() if r['tt'] == 'uppercase' else r['t']
            f = run.font
            naam, vet = lettertype(r)
            f.name, f.size, f.bold, f.italic = naam, Pt(r['s'] / 2), vet, r['i']
            if r['c']:
                f.color.rgb = rgb(r['c'])
                alfa(run._r.find(qn('a:rPr')).find(qn('a:solidFill')).find(qn('a:srgbClr')), r['c']['a'])
            else:
                # doorzichtige tekst: alleen een omlijning (-webkit-text-stroke), of onzichtbaar
                rPr = run._r.get_or_add_rPr()
                if r.get('sw') and r.get('sc'):
                    ln = rPr.makeelement(qn('a:ln'), {'w': str(int(Pt(r['sw'] / 2)))})
                    sf = ln.makeelement(qn('a:solidFill'), {})
                    sf.append(sf.makeelement(qn('a:srgbClr'), {'val': r['sc']['hex'].upper()}))
                    ln.append(sf)
                    rPr.insert(0, ln)
                rPr.insert(1 if r.get('sw') and r.get('sc') else 0, rPr.makeelement(qn('a:noFill'), {}))
            if r['ls']:
                run._r.get_or_add_rPr().set('spc', str(int(round(r['ls'] * 50))))


def tekstvak(slide, it):
    lh = it.get('lh')
    grootte = max((r['s'] for r in it['runs']), default=24)
    regel = lh or grootte * 1.3
    een_regel = it['h'] < regel * 1.6
    # PowerPoint zet dezelfde letter iets breder dan de browser: extra breedte, anders breekt een regel te vroeg af
    extra = 8 if een_regel else max(8, it['w'] * 0.02)
    x, w = it['x'], it['w'] + extra
    if it['al'] in ('right', 'end'):
        x -= extra
    elif it['al'] == 'center':
        x -= extra / 2
    tb = slide.shapes.add_textbox(e(x), e(it['y']), e(w), e(max(it['h'], regel)))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = not een_regel
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    vul_tekst(tf, it['runs'], it['al'], lh)
    return tb


# ------------------------------------------------------------------ vormen
def vak(slide, it):
    x, y, w, h = it['x'], it['y'], it['w'], it['h']
    hoek = 0
    if it.get('rot'):
        m = re.match(r'matrix\(([^,]+),\s*([^,]+)', it['rot'])
        if m:
            hoek = math.degrees(math.atan2(float(m.group(2)), float(m.group(1))))
            cx, cy = x + w / 2, y + h / 2
            w, h = it['ow'], it['oh']
            x, y = cx - w / 2, cy - h / 2
    bw, bc = it['bw'], it['bc']
    r = it['rad']
    uniform = all(b == bw[0] for b in bw) and all(c == bc[0] for c in bc) and bw[0] > 0
    if it['bg'] or uniform:
        if r[0] > 0 and r[0] == r[1] and r[2] == 0 and r[3] == 0:
            soort = MSO_SHAPE.ROUND_2_SAME_RECTANGLE
        elif max(r) > 0:
            soort = MSO_SHAPE.ROUNDED_RECTANGLE
        else:
            soort = MSO_SHAPE.RECTANGLE
        v = slide.shapes.add_shape(soort, e(x), e(y), e(w), e(h))
        v.shadow.inherit = False
        if soort != MSO_SHAPE.RECTANGLE:
            v.adjustments[0] = min(0.5, max(r) / max(1, min(w, h)))
            if soort == MSO_SHAPE.ROUND_2_SAME_RECTANGLE:
                v.adjustments[1] = 0
        vul(v, it['bg'], it.get('op', 1))
        lijn(v, bc[0] if uniform else None, bw[0] if uniform else 0)
        if hoek:
            v.rotation = hoek
        v.text_frame.text = ''
    if not uniform:
        # losse randen (scheidingslijnen, accentbalkjes) als dunne rechthoeken
        for i, (b, c) in enumerate(zip(bw, bc)):
            if b <= 0 or not c:
                continue
            rx, ry, rw, rh = [(x, y, w, b), (x + w - b, y, b, h), (x, y + h - b, w, b), (x, y, b, h)][i]
            lv = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, e(rx), e(ry), e(rw), e(rh))
            lv.shadow.inherit = False
            vul(lv, c)
            lijn(lv)


def verbinding(slide, it):
    if it.get('x1') is not None:
        x1, y1, x2, y2 = (float(it[k]) for k in ('x1', 'y1', 'x2', 'y2'))
        soort = MSO_CONNECTOR.ELBOW if it.get('route') == 'elbow' else MSO_CONNECTOR.STRAIGHT
    else:
        x1 = it['x']
        x2 = it['x'] + it['w']
        y1 = y2 = it['y'] + it['h'] / 2
        soort = MSO_CONNECTOR.STRAIGHT
    c = slide.shapes.add_connector(soort, e(x1), e(y1), e(x2), e(y2))
    if it['c']:
        c.line.color.rgb = rgb(it['c'])
    c.line.width = Pt(max(it['lw'], 2) / 2)
    ln = c.line._get_or_add_ln()
    kop = it.get('head') or 'end'
    if kop in ('end', 'both'):
        ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    if kop in ('start', 'both'):
        ln.append(ln.makeelement(qn('a:headEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))


def beeld(slide, it):
    src = it['src']
    pad = Path(unquote(urlparse(src).path.lstrip('/'))) if src.startswith('file:') else None
    if not pad or not pad.exists():
        print('  beeld ontbreekt:', src[:80])
        return
    with Image.open(pad) as im:
        iw, ih = im.size
    x, y, w, h = it['x'], it['y'], it['w'], it['h']
    if it.get('fit') == 'contain':
        s = min(w / iw, h / ih)
        x, y, w, h = x + (w - iw * s) / 2, y + (h - ih * s) / 2, iw * s, ih * s
    slide.shapes.add_picture(str(pad), e(x), e(y), e(w), e(h))


def celranden(cel, k, b):
    tcPr = cel._tc.get_or_add_tcPr()
    for plaats, kant in enumerate(('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB')):   # randen horen vóór de vulling
        ln = tcPr.makeelement(qn(kant), {'w': str(int(Pt(max(b, 0.5) / 2))), 'cmpd': 'sng'})
        if k:
            sf = ln.makeelement(qn('a:solidFill'), {})
            sf.append(sf.makeelement(qn('a:srgbClr'), {'val': k['hex'].upper()}))
            ln.append(sf)
        else:
            ln.append(ln.makeelement(qn('a:noFill'), {}))
        tcPr.insert(plaats, ln)


def tabel(slide, it):
    rijen = it['rows']
    nk = max(sum(c['cs'] for c in r) for r in rijen)
    gf = slide.shapes.add_table(len(rijen), nk, e(it['x']), e(it['y']), e(it['w']), e(it['h']))
    t = gf.table
    tblPr = t._tbl.tblPr
    tblPr.set('firstRow', '0')
    tblPr.set('bandRow', '0')
    stijl = tblPr.find(qn('a:tableStyleId'))
    if stijl is None:
        stijl = tblPr.makeelement(qn('a:tableStyleId'), {})
        tblPr.append(stijl)
    stijl.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'          # geen stijl, geen raster
    breedtes = [c['w'] for c in rijen[0]] if len(rijen[0]) == nk else [it['w'] / nk] * nk
    for j, b in enumerate(breedtes):
        t.columns[j].width = e(b)
    for i, r in enumerate(rijen):
        t.rows[i].height = e(max(c['h'] for c in r))
        j = 0
        for c in r:
            cel = t.cell(i, j)
            if c['cs'] > 1:
                cel.merge(t.cell(i, j + c['cs'] - 1))
            if c['bg']:
                cel.fill.solid()
                cel.fill.fore_color.rgb = rgb(c['bg'])
            else:
                cel.fill.background()
            p = c['pad']
            cel.margin_top, cel.margin_right, cel.margin_bottom, cel.margin_left = e(p[0]), e(p[1]), e(p[2]), e(p[3])
            cel.vertical_anchor = MSO_ANCHOR.TOP if c['va'] == 'top' else MSO_ANCHOR.MIDDLE
            cel.text_frame.word_wrap = True
            vul_tekst(cel.text_frame, c['runs'], c['al'], c.get('lh'))
            celranden(cel, c['bc'] or it['bc'], c['bw'] or it['bw'])
            j += c['cs']


# ------------------------------------------------------------------ animatie en overgang
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'


def overgang(slide):
    slide._element.append(etree.fromstring(
        f'<p:transition xmlns:p="{P}" xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" '
        'spd="med" p14:dur="700"><p:fade/></p:transition>'))


def wipe(slide, vormen):
    """Laat elke grafiek bij het openen van de dia vanzelf binnenschuiven (wipe), de een na de ander."""
    if not vormen:
        return
    # richting -> (preset, subtype, filter, duur): wipe voor grafieken, vervagen voor labels bij een grafiek
    filters = {'onder': ('22', '4', 'wipe(up)', 1400), 'links': ('22', '8', 'wipe(right)', 1400), 'vervaag': ('10', '0', 'fade', 500)}
    n = [3]

    def nid():
        n[0] += 1
        return n[0]

    effecten = []
    for k, (vorm, richting) in enumerate(vormen):
        preset, sub, filt, duur = filters[richting]
        spid = vorm.shape_id
        a, b, c = nid(), nid(), nid()
        effecten.append(
            f'<p:par><p:cTn id="{a}" presetID="{preset}" presetClass="entr" presetSubtype="{sub}" fill="hold" grpId="0" '
            f'nodeType="{"afterEffect" if k else "afterEffect"}"><p:stCondLst><p:cond delay="{0 if k == 0 else 200}"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
            f'</p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="{filt}"><p:cBhvr><p:cTn id="{c}" dur="{duur}"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par>')
    groep = ''.join(f'<p:par><p:cTn id="{nid()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
                    f'<p:childTnLst>{ef}</p:childTnLst></p:cTn></p:par>' for ef in effecten)
    bld = ''.join(f'<p:bldGraphic spid="{v.shape_id}" grpId="0"><p:bldAsOne/></p:bldGraphic>' if getattr(v, 'has_chart', False)
                  else f'<p:bldP spid="{v.shape_id}" grpId="0" animBg="1"/>' for v, _ in vormen)
    xml = (f'<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
           f'<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0">'
           f'<p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>{groep}</p:childTnLst></p:cTn></p:par>'
           f'</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           f'</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>{bld}</p:bldLst></p:timing>')
    slide._element.append(etree.fromstring(xml))


# ------------------------------------------------------------------ bouwen
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(1920 * EMU), Emu(1080 * EMU)
leeg = prs.slide_layouts[6]
deck = json.loads((DECK / 'project' / 'deck.json').read_text(encoding='utf-8'))
prs.core_properties.title = deck.get('title', '')
for i, sid in enumerate(deck['order'], 1):
    naam = f'{i:02d}_{sid}'
    mp = WERK / 'meting' / f'{naam}.json'
    if not mp.exists():
        print('geen meting:', naam)
        continue
    m = json.loads(mp.read_text(encoding='utf-8'))
    s = prs.slides.add_slide(leeg)
    ag = WERK / 'achtergrond' / f'{naam}.png'
    if ag.exists():
        s.shapes.add_picture(str(ag), 0, 0, prs.slide_width, prs.slide_height)
    elif m['bg']:
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(m['bg'])
    anim = []
    for it in m['items']:
        soort = it['type']
        if soort == 'box':
            vak(s, it)
        elif soort == 'text':
            tekstvak(s, it)
        elif soort == 'img':
            beeld(s, it)
        elif soort == 'line':
            verbinding(s, it)
        elif soort == 'table':
            tabel(s, it)
        elif soort == 'embed':
            sleutel = f"{sid}_{it['k']}"
            fn = grafieken.REGISTER.get((SLEUTEL, sleutel))
            if fn:
                bron = (WERK / 'embeds' / f'{sleutel}.html').read_text(encoding='utf-8')
                res = fn(s, it, bron)
                anim.extend(res if isinstance(res, list) else [res])
            elif (WERK / 'gif' / f'{sleutel}.gif').exists():
                s.shapes.add_picture(str(WERK / 'gif' / f'{sleutel}.gif'), e(it['x']), e(it['y']), e(it['w']), e(it['h']))
            else:
                print('  geen grafiek of GIF voor', sleutel)
    if m.get('notes'):
        s.notes_slide.notes_text_frame.text = m['notes']
    overgang(s)
    wipe(s, anim)
    print(naam, len(m['items']), 'objecten', f'{len(anim)} grafiek(en)' if anim else '')
UIT.parent.mkdir(parents=True, exist_ok=True)
prs.save(UIT)
print('opgeslagen:', UIT)
